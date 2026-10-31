"""Department volunteer ledger. Integer hundredths of an hour, never elapsed time."""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
from uuid import uuid4
import db

START = '2026-10-30'

def parse_hours(value):
    try:
        n = Decimal(str(value))
        if not n.is_finite() or n < 0 or n > 10000 or n * 100 != (n * 100).to_integral_value():
            raise ValueError()
        return int(n * 100)
    except (InvalidOperation, ValueError):
        raise ValueError('小时数需为 0～10000 的数字，最多两位小数；0 表示明确不计时长。')

def hours(units):
    return format(Decimal(units) / 100, '.2f')

def entry_date(value=None):
    value = value or date.today().isoformat()
    try:
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value): raise ValueError()
        d = date.fromisoformat(value)
    except ValueError:
        raise ValueError('日期格式：YYYY-MM-DD，例如 2026-10-30。')
    if value < START or d > date.today():
        raise ValueError('归属日期不得早于 2026-10-30，也不能是未来日期。')
    return value

def month_range(value=None):
    value = value or date.today().strftime('%Y-%m')
    try:
        if not re.fullmatch(r'\d{4}-\d{2}', value): raise ValueError()
        first = date.fromisoformat(value+'-01')
        end = date(first.year+1,1,1) if first.month==12 else date(first.year,first.month+1,1)
    except ValueError:
        raise ValueError('月份格式：YYYY-MM，例如 2026-11。')
    if value < START[:7] or value > date.today().strftime('%Y-%m'):
        raise ValueError('可查询 2026-10 至当前月份。')
    return value,max(first.isoformat(),START),end.isoformat()

def migrate(connect=None, backup=None):
    connect = connect or db.connect
    backup = backup or db._backup_before_migration
    with connect() as c:
        if db._table_exists(c,'volunteer_entries'):
            _backfill_pending(c)
            return None
        backup_path=backup(c,'volunteer_v9')
        c.execute('BEGIN IMMEDIATE')
        c.execute('''CREATE TABLE IF NOT EXISTS volunteer_entries (
            id INTEGER PRIMARY KEY AUTOINCREMENT, workspace TEXT NOT NULL, task_id INTEGER NOT NULL,
            qq TEXT NOT NULL, units INTEGER CHECK(units >= 0), service_date TEXT NOT NULL,
            actor_qq TEXT NOT NULL, created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
            UNIQUE(workspace,task_id,qq))''')
        c.execute('''CREATE TABLE IF NOT EXISTS volunteer_changes (
            id INTEGER PRIMARY KEY AUTOINCREMENT, entry_id INTEGER NOT NULL REFERENCES volunteer_entries(id),
            workspace TEXT NOT NULL, old_units INTEGER, new_units INTEGER,
            old_date TEXT, new_date TEXT NOT NULL, actor_qq TEXT NOT NULL,
            reason TEXT NOT NULL, created_at TEXT NOT NULL)''')
        c.execute('CREATE INDEX IF NOT EXISTS idx_volunteer_month ON volunteer_entries(workspace,service_date,qq)')
        _backfill_pending(c)
        c.execute('''CREATE TABLE IF NOT EXISTS volunteer_managers (workspace TEXT NOT NULL, qq TEXT NOT NULL, appointed_by TEXT NOT NULL, appointed_at TEXT NOT NULL, PRIMARY KEY(workspace,qq))''')
        return backup_path

def _backfill_pending(c):
    # Existing eligible completed work is pending registration, never assigned guessed hours.
    c.execute('''INSERT OR IGNORE INTO volunteer_entries
        (workspace,task_id,qq,units,service_date,actor_qq,created_at,updated_at)
        SELECT workspace,task_id,qq,NULL,substr(completed_at,1,10),COALESCE(reviewer_qq,''),?,?
        FROM task_assignments WHERE status='completed' AND substr(completed_at,1,10)>=?
        AND substr(completed_at,1,10)<=?''',(db.now_text(),db.now_text(),START,date.today().isoformat()))


def _admin(c,workspace,qq):
    return c.execute("""SELECT 1 FROM users u WHERE u.workspace=? AND u.qq=? AND u.active=1
        AND (u.role='owner' OR EXISTS
        (SELECT 1 FROM volunteer_managers m WHERE m.workspace=u.workspace AND m.qq=u.qq))""",(workspace,str(qq))).fetchone()

def can_manage(workspace,qq):
    with db.connect() as c:return bool(_admin(c,workspace,qq))

def managers(workspace):
    with db.connect() as c:
        return [dict(r) for r in c.execute("""SELECT u.qq,u.name FROM volunteer_managers m
            JOIN users u ON u.workspace=m.workspace AND u.qq=m.qq WHERE m.workspace=? AND u.active=1 ORDER BY u.id""",(workspace,))]

def set_manager(workspace,actor,target,add):
    with db.connect() as c:
        c.execute('BEGIN IMMEDIATE')
        if not _admin(c,workspace,actor):
            raise ValueError('只有总负责人或高级管理员可以任命或移除志愿时长负责人。')
        if add and not c.execute('SELECT 1 FROM users WHERE workspace=? AND qq=? AND active=1',(workspace,str(target))).fetchone():
            raise ValueError('只能任命有效成员。')
        if add:
            c.execute('INSERT OR IGNORE INTO volunteer_managers(workspace,qq,appointed_by,appointed_at) VALUES(?,?,?,?)',(workspace,str(target),str(actor),db.now_text()))
        else:c.execute('DELETE FROM volunteer_managers WHERE workspace=? AND qq=?',(workspace,str(target)))
        c.execute('INSERT INTO audit_logs(workspace,actor_qq,action,detail,created_at) VALUES(?,?,?,?,?)',
                  (workspace,str(actor),'时长负责人添加' if add else '时长负责人移除',str(target),db.now_text()))

def _change(c,row,units,day,actor,reason):
    stamp=db.now_text()
    c.execute('''INSERT INTO volunteer_changes(entry_id,workspace,old_units,new_units,old_date,new_date,actor_qq,reason,created_at)
        VALUES(?,?,?,?,?,?,?,?,?)''',(row['id'],row['workspace'],row['units'],units,row['service_date'],day,str(actor),reason,stamp))
    c.execute('UPDATE volunteer_entries SET units=?,service_date=?,actor_qq=?,updated_at=? WHERE id=?',(units,day,str(actor),stamp,row['id']))
    c.execute('INSERT INTO audit_logs(workspace,actor_qq,action,detail,created_at) VALUES(?,?,?,?,?)',
        (row['workspace'],str(actor),'志愿时长登记',f"记录 #{row['id']}｜任务 #{row['task_id']}｜{hours(units)} 小时｜{day}｜{reason}",stamp))

def on_approval(c,workspace,task_id,qq,actor,units=None,day=None):
    day=entry_date(day)
    stamp=db.now_text()
    row=c.execute('SELECT * FROM volunteer_entries WHERE workspace=? AND task_id=? AND qq=?',(workspace,task_id,str(qq))).fetchone()
    if row and row['units'] is not None:
        if units is not None: raise ValueError('该任务成员已有时长记录；请不带时长审核，需调整时使用“更正志愿时长”，不会重复累计。')
        return
    if not row:
        c.execute('INSERT INTO volunteer_entries(workspace,task_id,qq,units,service_date,actor_qq,created_at,updated_at) VALUES(?,?,?,NULL,?,?,?,?)',
                  (workspace,task_id,str(qq),day,str(actor),stamp,stamp))
        row=c.execute('SELECT * FROM volunteer_entries WHERE workspace=? AND task_id=? AND qq=?',(workspace,task_id,str(qq))).fetchone()
    if units is not None: _change(c,row,units,day,actor,'审核通过时登记')

def register(workspace,task_id,qq,actor,value,day=None):
    units=parse_hours(value);day=entry_date(day)
    with db.connect() as c:
        c.execute('BEGIN IMMEDIATE')
        if not _admin(c,workspace,actor): raise ValueError('仅有效总负责人或志愿时长负责人可登记时长。')
        a=c.execute("SELECT completed_at FROM task_assignments WHERE workspace=? AND task_id=? AND qq=? AND status='completed'",(workspace,task_id,str(qq))).fetchone()
        if not a or not a['completed_at'] or a['completed_at'][:10]<START:
            raise ValueError('只能补录 2026-10-30 起审核完成的个人任务份额。')
        on_approval(c,workspace,task_id,qq,actor,units,day)

def correct(workspace,entry_id,actor,value,day,reason):
    units=parse_hours(value);day=entry_date(day);reason=reason.strip()
    if not reason or len(reason)>500: raise ValueError('更正必须填写原因，最多 500 字。')
    with db.connect() as c:
        c.execute('BEGIN IMMEDIATE')
        if not _admin(c,workspace,actor): raise ValueError('仅有效总负责人或志愿时长负责人可更正时长。')
        row=c.execute('SELECT * FROM volunteer_entries WHERE workspace=? AND id=?',(workspace,entry_id)).fetchone()
        if not row: raise ValueError('时长记录不存在。')
        _change(c,row,units,day,actor,reason)

def report(workspace,month=None,qq=None):
    month,start,end=month_range(month)
    with db.connect() as c:
        details=[dict(r) for r in c.execute('''SELECT e.*,COALESCE(u.name,e.qq) AS name,
            COALESCE(t.subject,'') AS subject FROM volunteer_entries e
            LEFT JOIN users u ON u.workspace=e.workspace AND u.qq=e.qq
            LEFT JOIN tasks t ON t.workspace=e.workspace AND t.id=e.task_id
            WHERE e.workspace=? AND e.service_date>=? AND e.service_date<?
            AND (? IS NULL OR e.qq=?) ORDER BY e.service_date,e.id''',(workspace,start,end,qq,qq))]
        people=[dict(r) for r in c.execute('''SELECT qq,name,student_number,major_college,minor_college,active
            FROM users WHERE workspace=? AND (? IS NULL OR qq=?) AND (active=1 OR qq IN
            (SELECT qq FROM volunteer_entries WHERE workspace=? AND service_date>=? AND service_date<?)) ORDER BY id''',
            (workspace,qq,qq,workspace,start,end))]
    sums={}
    for r in details:
        v=sums.setdefault(r['qq'],[0,0,0]);v[0]+=r['units'] or 0;v[1]+=r['units'] is None;v[2]+=1
    for p in people:
        p['units'],p['pending'],p['entries']=sums.get(p['qq'],[0,0,0])
    return month,start,end,people,details

def export_month(workspace,month,base):
    from openpyxl import Workbook
    m,start,end,people,details=report(workspace,month)
    folder=Path(base)/'exports';folder.mkdir(exist_ok=True)
    path=folder/f'volunteer_{m}_{uuid4().hex[:8]}.xlsx'
    wb=Workbook();summary=wb.active;summary.title='人员汇总'
    summary.append(['统计月份','起始日期（含）','截止日期（不含）']);summary.append([m,start,end]);summary.append([])
    summary.append(['姓名','QQ','学号','主修学院','辅修学院','成员有效','已登记小时','待登记份数'])
    def safe(row):
        # Force all strings to text, including names beginning with '='.
        sheet,row_values=row
        sheet.append(row_values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    for p in people:safe((summary,[p['name'],p['qq'],p['student_number'],p['major_college'],p['minor_college'],p['active'],p['units']/100,p['pending']]))
    d=wb.create_sheet('登记明细');d.append(['记录编号','姓名','QQ','任务编号','学科','归属日期','小时数','状态','登记人','创建时间','更新时间'])
    for r in details:safe((d,[r['id'],r['name'],r['qq'],r['task_id'],r['subject'],r['service_date'],None if r['units'] is None else r['units']/100,'待登记' if r['units'] is None else '已登记',r['actor_qq'],r['created_at'],r['updated_at']]))
    changes=wb.create_sheet('更正留痕');changes.append(['记录编号','原小时数','新小时数','原日期','新日期','操作人','原因','操作时间'])
    with db.connect() as c:
        rows=c.execute('''SELECT * FROM volunteer_changes WHERE workspace=? AND
            ((old_date>=? AND old_date<?) OR (new_date>=? AND new_date<?)) ORDER BY id''',(workspace,start,end,start,end)).fetchall()
    for r in rows:safe((changes,[r['entry_id'],None if r['old_units'] is None else r['old_units']/100,r['new_units']/100,r['old_date'],r['new_date'],r['actor_qq'],r['reason'],r['created_at']]))
    for sheet in wb:
        sheet.freeze_panes='A2'
        for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=22
    wb.save(path);wb.close();return path


def has_appointment(workspace,qq):
    # Include inactive appointments so a normal admin cannot restore a senior user.
    with db.connect() as c:
        return bool(c.execute('SELECT 1 FROM volunteer_managers WHERE workspace=? AND qq=?',(workspace,str(qq))).fetchone())
