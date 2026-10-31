import asyncio,tempfile,zipfile
from pathlib import Path
from openpyxl import load_workbook
import bot,db,volunteer as v
class Client:
 def __init__(self):self.messages=[]
 async def send_private(self,q,t):self.messages.append((str(q),t))
 async def send_group(self,*a):pass
async def checks(tmp):
 app=bot.NotesDepartmentBot({'department_group_id':'w','owner_qq':'owner','owner_name':'负责人'});app.client=Client()
 async def send(s,q='owner'):
  await app.handle_private({'user_id':q,'raw_message':s})
  return app.client.messages[-1][1]
 await send('时长负责人添加 manager','admin');assert not v.can_manage('w','manager')
 await send('时长负责人添加 manager');assert v.can_manage('w','manager')
 assert bot.is_admin(db.get_user('w','manager')) and not v.can_manage('w','admin')
 assert 'manager' in [r['qq'] for r in db.list_admins('w')]
 assert '高级管理员' in await send('我的身份','manager')
 for q in ('m','admin','manager','owner'):
  text=await send('全部帮助',q)
  assert ('发布任务' in text and '所有任务' in text)==(q!='m')
  assert ('更正志愿时长' in text)==(q in ('manager','owner'))
 for cmd in bot.SENIOR_COMMANDS:
  for q in ('m','admin'):
   assert '仅总负责人和高级管理员' in await send(cmd,q)
 for cmd in bot.ADMIN_COMMANDS:
  assert '仅管理者' in await send(cmd,'m')
 for cmd in ('任务','帮助','发布任务','我的志愿时长','乱写'):
  assert '还没有注册' in await send(cmd,'unknown')
 for cmd in ('管理员移除 manager','成员移除 manager','管理员添加 manager','成员移除 owner'):
  await send(cmd,'admin')
 assert db.get_user('w','manager') and v.can_manage('w','manager')
 await send('时长负责人添加 second','manager');assert v.can_manage('w','second')
 await send('时长负责人移除 second','manager');assert not v.can_manage('w','second')
 t=db.create_task('w','笔记','数学','集合',1,'owner');db.claim_task('w',t,'m');db.submit_task('w',t,'m','成果')
 await send(f'通过 {t} m 1.50','admin');assert db.get_assignment('w',t,'m')['status']=='submitted'
 try:db.approve_submission('w',t,'m','admin','1.50')
 except ValueError:pass
 else:raise AssertionError('database permission bypass')
 await send(f'通过 {t} m','admin');assert db.get_assignment('w',t,'m')['status']=='completed'
 await send(f'登记志愿时长 {t} m 1.50','manager')
 assert '1.50' in await send('我的志愿时长','m')
 assert '1.50' in await send('志愿时长统计','manager')
 rid=db.create_task_request('w','m',['笔记','数学','集合','1','补齐'])
 assert db.decide_task_request('w',rid,'manager','approved')[0]=='ok'
 # Senior also has task-review rights, independent of stored member role.
 t2=db.create_task('w','笔记','数学','集合',1,'owner');db.claim_task('w',t2,'m');db.submit_task('w',t2,'m','成果')
 app.client.messages=[];await send(f'通过 {t2} m 2','manager')
 assert db.get_assignment('w',t2,'m')['status']=='completed'
 for q in ('owner','manager'):assert len([x for x in app.client.messages if x[0]==q and '【志愿时长提醒】' in x[1]])==1
 db.log_action('w','owner','普通操作','允许')
 db.log_action('w','owner','志愿时长更正','秘密时长')
 assert db.list_audit_logs('w',1,False)[0]['action']=='普通操作'
 for include in (True,False):
  path=bot.make_export_xlsx('w',include);wb=load_workbook(path)
  assert ('志愿时长' in wb.sheetnames)==include
  assert ('秘密时长' in str(list(wb['操作日志'].values)))==include
  wb.close()
  with zipfile.ZipFile(bot.make_export_csv_zip('w',include)) as z:
   assert any(n.startswith('volunteer_') for n in z.namelist())==include
   assert ('秘密时长' in z.read('audit_logs.csv').decode('utf-8-sig'))==include
 await send('发布任务','admin');assert 'admin' in bot.publish_sessions
 db.set_role('w','admin','member');await send('笔记','admin');assert 'admin' not in bot.publish_sessions
 with db.connect() as c:c.execute("UPDATE users SET active=0 WHERE qq='manager'")
 bot.profile_sessions['manager']=[]
 assert '停用' in await send('注册 manager','manager')
 assert 'manager' not in bot.profile_sessions and not v.can_manage('w','manager')
 db.set_role('w','admin','admin');await send('成员恢复 manager','admin')
 assert not db.get_user('w','manager')
 await send('成员恢复 manager');assert v.can_manage('w','manager')
 await send('时长负责人移除 manager');assert not bot.is_admin(db.get_user('w','manager'))
with tempfile.TemporaryDirectory() as tmp:
 tmp=Path(tmp);db.DB_PATH=tmp/'roles.db';bot.BASE_DIR=tmp;db.init_db();db.ensure_owner('w','owner')
 for q in ('admin','manager','m','second'):db.register_user('w',q,q)
 db.set_role('w','admin','admin')
 asyncio.run(checks(tmp))
 print('PASS: role menus/command gates; unknown/inactive users; senior appointment/review; ordinary admin no ledger access; own hours; export/audit isolation; owner protection; revoked sessions. No network messages.')
