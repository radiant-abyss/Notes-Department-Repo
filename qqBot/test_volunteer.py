import asyncio,sqlite3,tempfile
from contextlib import closing
from pathlib import Path
from datetime import date
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor
from openpyxl import load_workbook
import db,bot,volunteer as v
class FutureDate(date):
 @classmethod
 def today(cls):return cls(2026,12,15)
class Client:
 def __init__(self):self.messages=[];self.fail=set()
 async def send_private(self,qq,text):
  self.messages.append((str(qq),text))
  if str(qq) in self.fail:raise ConnectionError('test failure')
 async def send_group(self,*args):pass
 async def send_private_file(self,*args):raise ConnectionError('test file failure')
def completed_task(member='m',hours=None):
 t=db.create_task('w','笔记','数学','集合',1,'admin')
 assert db.claim_task('w',t,member)=='ok'
 assert db.submit_task('w',t,member,'成果')=='ok'
 assert db.approve_submission('w',t,member,'admin',hours)=='ok'
 return t
async def commands(tmp):
 app=bot.NotesDepartmentBot({'department_group_id':'w','owner_qq':'owner','owner_name':'负责人'});app.client=Client()
 async def send(text,q='m'):await app.handle_private({'user_id':q,'raw_message':text})
 t=db.create_task('w','笔记','数学','集合',1,'admin');db.claim_task('w',t,'m')
 await send(f'提交 {t}');app.client.fail={'owner','m'};await send('成果链接')
 assert len([x for x in app.client.messages if x[0]=='owner' and '新的任务提交' in x[1]])==1
 assert any(x[0]=='admin' and '新的任务提交' in x[1] for x in app.client.messages)
 assert db.get_assignment('w',t,'m')['status']=='submitted'
 app.client.fail=set();await send(f'通过 {t} m 2.5','owner')
 assert db.get_assignment('w',t,'m')['status']=='completed'
 await send('我的志愿时长');assert '2.50' in app.client.messages[-1][1]
 await send('志愿时长统计');assert '高级管理员' in app.client.messages[-1][1]
 await send('导出志愿时长','owner');assert '月报已生成，但发送失败' in app.client.messages[-1][1]
 # Same owner/publisher is notified once.
 t2=db.create_task('w','笔记','数学','集合',1,'owner');db.claim_task('w',t2,'m')
 app.client.messages=[];await send(f'提交 {t2}');await send('第二份成果')
 assert len([x for x in app.client.messages if x[0]=='owner' and '新的任务提交' in x[1]])==1
 await send(f'通过 {t2} m bad','owner');assert db.get_assignment('w',t2,'m')['status']=='submitted'
 await send(f'通过 {t2} m','owner');assert '待登记' in app.client.messages[-2][1] or any('待登记' in x[1] for x in app.client.messages)
 await send(f'登记志愿时长 {t2} m 1.25 2026-10-30','owner')
 assert app.client.messages[-1][1].startswith('志愿时长已登记')
 await send('志愿时长明细 2026-10');assert '记录 #' in app.client.messages[-1][1]
 await send('我的志愿时长 2026-13');assert '月份格式' in app.client.messages[-1][1]
with tempfile.TemporaryDirectory() as tmp:
 tmp=Path(tmp);db.DB_PATH=tmp/'test.db';bot.BASE_DIR=tmp;db.init_db();db.ensure_owner('w','owner')
 for q in ('admin','m','zero'):db.register_user('w',q,'=FORMULA' if q=='zero' else q)
 with db.connect() as c:c.execute("UPDATE users SET role='admin' WHERE qq='admin'")
 v.set_manager('w','owner','admin',True)
 assert v.migrate() is None
 for x in ('NaN','Infinity','-1','0.001','text'):
  try:v.parse_hours(x)
  except ValueError:pass
  else:raise AssertionError(x)
 assert v.parse_hours('2.50')==250 and v.parse_hours('0')==0
 with patch.object(v,'date',FutureDate):
  t=completed_task(hours='1.25');assert db.approve_submission('w',t,'m','admin','1.25')=='not_submitted'
  rows=v.report('w','2026-12')[4];assert len(rows)==1 and rows[0]['units']==125
  eid=rows[0]['id'];v.correct('w',eid,'owner','2.50','2026-11-30','调整工作量及月份')
  assert not v.report('w','2026-12')[4] and v.report('w','2026-11')[4][0]['units']==250
  try:v.correct('w',eid,'m','99','2026-11-30','越权')
  except ValueError:pass
  else:raise AssertionError('permission')
  try:v.correct('other',eid,'owner','99','2026-11-30','跨群')
  except ValueError:pass
  else:raise AssertionError('workspace isolation')
  assert db.reopen_task('w',t);db.submit_task('w',t,'m','修订')
  try:db.approve_submission('w',t,'m','admin','3')
  except ValueError:pass
  else:raise AssertionError('duplicate ledger')
  assert db.get_assignment('w',t,'m')['status']=='submitted'
  assert db.approve_submission('w',t,'m','admin')=='ok'
  assert v.report('w','2026-11')[4][0]['units']==250
  missing=completed_task();assert v.report('w','2026-12')[3][2]['pending']==1
  with ThreadPoolExecutor(2) as pool:
   def register(_):
    try:v.register('w',missing,'m','admin','1','2026-10-30');return 'ok'
    except ValueError:return 'duplicate'
   assert sorted(pool.map(register,range(2)))==['duplicate','ok']
  old=completed_task()
  with db.connect() as c:
   c.execute("UPDATE task_assignments SET completed_at='2026-10-29 12:00:00' WHERE task_id=?",(old,))
   c.execute('DELETE FROM volunteer_entries WHERE task_id=?',(old,))
  try:v.register('w',old,'m','admin','1','2026-10-30')
  except ValueError:pass
  else:raise AssertionError('cutoff')
  for d in ('2026-10-29','2026-12-16','2026-02-30'):
   try:v.entry_date(d)
   except ValueError:pass
   else:raise AssertionError(d)
  v.correct('w',eid,'owner','0','2026-11-30','冲销')
  assert v.report('w','2026-11')[4][0]['units']==0
  with db.connect() as c:c.execute("UPDATE users SET active=0 WHERE qq='m'")
  assert any(p['qq']=='m' for p in v.report('w','2026-11')[3])
  path=v.export_month('w','2026-11',tmp)
  wb=load_workbook(path);assert len(wb.sheetnames)==3
  assert any(c.value=='=FORMULA' and c.data_type=='s' for row in wb['人员汇总'] for c in row);wb.close()
  with db.connect() as c:c.execute("UPDATE users SET active=1 WHERE qq='m'")
 # Use a fresh workspace dataset for command assertions.
 db.DB_PATH=tmp/'commands.db';db.init_db();db.ensure_owner('w','owner')
 for q in ('admin','m'):db.register_user('w',q,q)
 with db.connect() as c:c.execute("UPDATE users SET role='admin' WHERE qq='admin'")
 asyncio.run(commands(tmp))
 print('PASS: exact hour units; month/date cutoff; missing vs zero; atomic approval; concurrent registration; correction/moving month/audit; reopen no double credit; inactive history; permissions; safe Excel; owner always attempted/dedup; notification failures; private commands. No real network messages.')
